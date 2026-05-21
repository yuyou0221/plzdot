import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUMMARY_COLUMNS = [
    ("projectId", "项目编号"),
    ("projectName", "项目名称"),
    ("team", "所属团队"),
    ("owner", "项目管理"),
    ("projectStartDate", "项目启动日期"),
    ("effectiveProjectStartDate", "有效项目启动日期"),
    ("plannedProjectStartDate", "计划应启动日期"),
    ("projectStartDeltaDays", "启动偏差天数"),
    ("plannedLaunchDate", "计划出货日期"),
    ("plannedProductionMilestone", "正推计划出货节点"),
    ("plannedTotalScheduleDays", "项目完整排期天数"),
    ("projectedLaunchDate", "当前测算出货日期"),
    ("launchDeltaDays", "出货偏差天数"),
    ("projectFeasibility", "计划可行性"),
    ("route", "路线"),
    ("hasThreeView", "是否需要三视图"),
    ("status", "项目阶段"),
]

TASK_COLUMNS = [
    ("projectId", "项目编号"),
    ("projectName", "项目名称"),
    ("projectStatus", "项目阶段"),
    ("plannedLaunchDate", "计划出货日期"),
    ("projectedLaunchDate", "当前测算出货日期"),
    ("launchDeltaDays", "出货偏差天数"),
    ("taskId", "任务ID"),
    ("taskName", "任务名称"),
    ("durationDays", "标准工期"),
    ("taskStatus", "任务状态"),
    ("autoStarted", "是否当前应开始"),
    ("missingActualPredecessorIds", "缺少实际完成的前置任务"),
    ("actualStartDate", "实际开始日期"),
    ("actualFinishDate", "实际完成日期"),
    ("expectedFinishDate", "推进中任务预期完成时间"),
    ("inferredCompleted", "是否推断完成"),
    ("inferredCompletionDate", "推断完成参考日期"),
    ("plannedStartDate", "正常节奏开始日"),
    ("plannedFinishDate", "正常节奏完成日"),
    ("forecastStartDate", "根据当前进度预测开始日"),
    ("forecastFinishDate", "根据当前进度预测完成日"),
    ("calculatedStartDate", "当前测算开始日"),
    ("calculatedFinishDate", "当前测算完成日"),
    ("currentDdlDate", "当前DDL日期"),
    ("originalLatestStartDate", "原计划上线最晚开始日"),
    ("originalLatestFinishDate", "原计划上线最晚完成日"),
    ("latestStartDate", "再次延期警告最晚开始日"),
    ("latestFinishDate", "再次延期警告最晚完成日"),
    ("floatDays", "安全缓冲天数"),
    ("planDeltaDays", "相对正常节奏偏差天数"),
    ("deadlineRiskDays", "再次延期风险天数"),
    ("warningWindowDays", "距离再次延期剩余天数"),
    ("impactStatus", "影响状态"),
    ("riskLevel", "风险等级"),
    ("isBlockingLaunch", "是否阻塞出货"),
]

BASIS_LABELS = {
    "actual_finish": "实际完成日期",
    "remaining_days": "剩余工期",
    "actual_start_plus_duration": "实际开始日期+标准工期",
    "downstream_inferred": "后置任务反推完成",
    "baseline": "正常节奏基线",
}


def localized(value):
    if value is True:
        return "是"
    if value is False:
        return "否"
    if value in BASIS_LABELS:
        return BASIS_LABELS[value]
    return "" if value is None else value


def write_sheet(wb, title, rows, columns):
    ws = wb.create_sheet(title)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([localized(row.get(key)) for key, _ in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, (_, label) in enumerate(columns, start=1):
        values = [label]
        for row in rows[:200]:
            values.append(str(localized(row.get(columns[col_idx - 1][0]))))
        width = min(max(max(len(v) for v in values) + 2, 10), 34)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def main():
    if len(sys.argv) < 3:
        raise SystemExit("Usage: export_chinese_excel.py <input.json> <output.xlsx>")

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    data = json.loads(input_path.read_text(encoding="utf-8"))

    wb = Workbook()
    wb.remove(wb.active)

    write_sheet(wb, "项目摘要", data.get("projects", []), SUMMARY_COLUMNS)
    write_sheet(wb, "后续任务测算", data.get("futureRows", []), TASK_COLUMNS)
    write_sheet(wb, "全部任务明细", data.get("rows", []), TASK_COLUMNS)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(str(output_path))


if __name__ == "__main__":
    main()
